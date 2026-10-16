import os
import sys
import re
from pathlib import Path
import docx
import pypdf

SOURCE_DIR = Path(r"C:\Projects\allchildcaresolutions\2026-09-10-forms\raw_unpacked")
OUT_FORMS = Path(r"C:\Projects\allchildcaresolutions\docs\forms_source_of_truth")
OUT_POLICIES = Path(r"C:\Projects\allchildcaresolutions\docs\policies_source_of_truth")

OUT_FORMS.mkdir(parents=True, exist_ok=True)
OUT_POLICIES.mkdir(parents=True, exist_ok=True)

def extract_docx(file_path: Path) -> str:
    try:
        doc = docx.Document(file_path)
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip().replace('\n', ' ') for c in row.cells if c.text.strip()]
                if row_cells:
                    full_text.append(" | ".join(row_cells))
        return "\n\n".join(full_text)
    except Exception as e:
        return f"ERROR EXTRACTING DOCX: {e}"

def extract_pdf(file_path: Path) -> str:
    try:
        reader = pypdf.PdfReader(file_path)
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages_text.append(f"<!-- Page {i+1} -->\n" + text.strip())
        return "\n\n".join(pages_text)
    except Exception as e:
        return f"ERROR EXTRACTING PDF: {e}"

def sanitize_filename(name: str) -> str:
    clean = re.sub(r'[^\w\s-]', '', name).strip().lower()
    return re.sub(r'[-\s]+', '-', clean)

def main():
    print(f"Scanning {SOURCE_DIR}...")
    forms_inventory = []
    policies_inventory = []

    for root, dirs, files in os.walk(SOURCE_DIR):
        for f in files:
            p = Path(root) / f
            ext = p.suffix.lower()
            rel = p.relative_to(SOURCE_DIR)
            is_policy = "policy" in str(rel).lower()
            
            if ext in ['.docx', '.pdf', '.txt']:
                content = ""
                if ext == '.docx':
                    content = extract_docx(p)
                elif ext == '.pdf':
                    content = extract_pdf(p)
                elif ext == '.txt':
                    content = p.read_text(encoding='utf-8', errors='ignore')

                dest_root = OUT_POLICIES if is_policy else OUT_FORMS
                
                # Determine subcategory folder
                parts = rel.parts
                # e.g. ('forms', 'FINALLLLLLL FORMS', 'Subdir', 'file.pdf')
                # skip first 2 parts if they are just wrapper names
                meaningful_parts = [part for part in parts[:-1] if not part.startswith('FINAL') and not part.startswith('xACCS') and part not in ['forms', 'policy']]
                
                target_folder = dest_root
                if meaningful_parts:
                    target_folder = dest_root / Path(*meaningful_parts)
                target_folder.mkdir(parents=True, exist_ok=True)
                
                md_filename = p.stem + ".md"
                md_path = target_folder / md_filename
                
                header = f"# {p.stem}\n\n"
                header += f"- **Original Source File**: `{p.name}`\n"
                header += f"- **Format**: `{ext.upper()}`\n"
                header += f"- **Original Relative Path**: `{rel}`\n\n---\n\n"
                
                md_path.write_text(header + content, encoding='utf-8')
                
                item_info = {
                    "name": p.name,
                    "stem": p.stem,
                    "ext": ext,
                    "md_path": str(md_path),
                    "size": p.stat().st_size,
                    "is_policy": is_policy,
                    "category": "/".join(meaningful_parts) if meaningful_parts else "General"
                }
                
                if is_policy:
                    policies_inventory.append(item_info)
                else:
                    forms_inventory.append(item_info)

    print(f"Processed {len(forms_inventory)} form documents into {OUT_FORMS}")
    print(f"Processed {len(policies_inventory)} policy documents into {OUT_POLICIES}")

    # Generate a comprehensive Markdown index
    index_md = "# ACCS Source of Truth: Master Document Index & Pipeline Audit\n\n"
    index_md += f"Generated from extracted archives on 2026-09-10.\n\n"
    
    index_md += f"## 1. Forms Inventory ({len(forms_inventory)} documents)\n\n"
    index_md += "| Category | Document Name | Source Format | Extracted Markdown |\n"
    index_md += "|---|---|---|---|\n"
    for item in sorted(forms_inventory, key=lambda x: (x['category'], x['name'])):
        index_md += f"| {item['category']} | {item['name']} | `{item['ext']}` | [{item['stem']}]({item['md_path']}) |\n"
        
    index_md += f"\n## 2. Policies Inventory ({len(policies_inventory)} documents)\n\n"
    index_md += "| Category | Policy Document Name | Source Format | Extracted Markdown |\n"
    index_md += "|---|---|---|---|\n"
    for item in sorted(policies_inventory, key=lambda x: (x['category'], x['name'])):
        index_md += f"| {item['category']} | {item['name']} | `{item['ext']}` | [{item['stem']}]({item['md_path']}) |\n"

    index_path = Path(r"C:\Projects\allchildcaresolutions\docs\MASTER_FORMS_AND_POLICIES_INDEX.md")
    index_path.write_text(index_md, encoding='utf-8')
    print(f"Master index written to {index_path}")

if __name__ == "__main__":
    main()
